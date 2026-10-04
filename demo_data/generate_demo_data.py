import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from docx import Document

DEMO_DIR = Path(__file__).resolve().parent

def generate_pdf():
    pdf_path = DEMO_DIR / "1_Supply_Contract.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    width, height = letter

    # Page 1: Parties, Purpose, Value, Payment Terms
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, height - 50, "COMMERCIAL SUPPLY AGREEMENT")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 70, "Date of Execution: 10 Jan 2026 | Agreement Ref: OS-NW-2026-004")
    c.line(50, height - 75, width - 50, height - 75)

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 100, "1. PARTIES")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 120, "Buyer: Orion Interiors, 402 Business Park, Mumbai, Maharashtra.")
    c.drawString(50, height - 135, "Seller: Northwind Supplies Pvt Ltd, 12 Industrial Area, Gurugram, Haryana.")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 165, "2. SCOPE AND COMMERCIAL TERMS")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 185, "Clause 2.1 - Contract Value: The total contract value is agreed at INR 5,00,000 plus 18% GST")
    c.drawString(50, height - 200, "(Net Payable: INR 5,90,000).")
    c.drawString(50, height - 220, "Clause 2.2 - Advance Payment: Buyer shall pay a 50% advance amounting to INR 2,50,000")
    c.drawString(50, height - 235, "upon signing of this agreement before manufacturing commences.")
    c.drawString(50, height - 255, "Clause 2.3 - Balance Payment: Balance INR 3,40,000 shall be payable upon final delivery")
    c.drawString(50, height - 270, "and satisfactory inspection at buyer's facility.")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 305, "3. DELIVERY SCHEDULE AND LIQUIDATED DAMAGES")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 325, "Clause 3.1 - Delivery Deadline: Seller strictly covenants that full delivery shall be completed")
    c.drawString(50, height - 340, "on or before 15 Feb 2026.")
    c.drawString(50, height - 360, "Clause 3.2 - Late Penalty: In the event of delayed delivery, a late penalty of 1% of the order value")
    c.drawString(50, height - 375, "per week of delay shall be deducted from the final invoice payment.")

    c.setFont("Helvetica-Oblique", 9)
    c.drawString(50, 40, "Page 1 of 4 - Commercial Supply Agreement (Orion Interiors & Northwind Supplies)")
    c.showPage()

    # Page 2: Specifications and Quality
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, height - 50, "COMMERCIAL SUPPLY AGREEMENT (Contd.)")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 80, "4. SPECIFICATIONS AND MATERIALS")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 100, "All ergonomic executive desks, conference tables, and mesh task chairs must conform")
    c.drawString(50, height - 115, "to Grade-A commercial architectural grade ISO-9001 quality benchmarks.")
    c.drawString(50, height - 135, "All finishes must be matte walnut laminate with powder-coated steel understructures.")
    c.drawString(50, height - 160, "Any deviation requires prior written sign-off from Orion Interiors project director.")

    c.setFont("Helvetica-Oblique", 9)
    c.drawString(50, 40, "Page 2 of 4 - Commercial Supply Agreement (Orion Interiors & Northwind Supplies)")
    c.showPage()

    # Page 3: Inspection, Termination and Dispute Resolution
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, height - 50, "COMMERCIAL SUPPLY AGREEMENT (Contd.)")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 80, "5. INSPECTION AND DAMAGED GOODS")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 100, "Buyer reserves the right to inspect delivered consignments within 7 business days.")
    c.drawString(50, height - 115, "Damaged or sub-standard articles shall be rectified or replaced within 10 days by Seller.")

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, height - 145, "6. TERMINATION AND JURISDICTION")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 165, "This agreement is governed by the laws of India, subject to the jurisdiction of Mumbai courts.")
    c.drawString(50, height - 180, "Note: Intentionally no warranty clause is included in this agreement.")

    c.setFont("Helvetica-Oblique", 9)
    c.drawString(50, 40, "Page 3 of 4 - Commercial Supply Agreement (Orion Interiors & Northwind Supplies)")
    c.showPage()

    # Page 4: Signatures and Execution
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, height - 50, "COMMERCIAL SUPPLY AGREEMENT - SIGNATURES")
    c.setFont("Helvetica", 10)
    c.drawString(50, height - 80, "IN WITNESS WHEREOF, the authorized representatives have executed this agreement on 10 Jan 2026.")
    
    c.drawString(50, height - 130, "For Orion Interiors (Buyer):")
    c.drawString(50, height - 150, "Signature: [Signed - Rajesh Mehta, Managing Director]")
    c.drawString(50, height - 165, "Date: 10 Jan 2026")

    c.drawString(320, height - 130, "For Northwind Supplies Pvt Ltd (Seller):")
    c.drawString(320, height - 150, "Signature: [Signed - Vikram Sharma, Director]")
    c.drawString(320, height - 165, "Date: 10 Jan 2026")

    c.setFont("Helvetica-Oblique", 9)
    c.drawString(50, 40, "Page 4 of 4 - Commercial Supply Agreement (Orion Interiors & Northwind Supplies)")
    c.showPage()

    c.save()
    print("Generated 1_Supply_Contract.pdf")

def generate_invoice_image():
    # 2_Invoice.jpg (Scanned photo slightly skewed)
    img_path = DEMO_DIR / "2_Invoice.jpg"
    w, h = 1000, 1300
    img = Image.new("RGB", (w, h), color=(250, 249, 245))
    draw = ImageDraw.Draw(img)

    # Header
    draw.rectangle([(40, 40), (960, 140)], fill=(235, 238, 242), outline=(180, 185, 195))
    draw.text((60, 55), "NORTHWIND SUPPLIES PVT LTD", fill=(20, 20, 30))
    draw.text((60, 85), "TAX INVOICE | GSTIN: 07AAAAA0000A1Z5", fill=(60, 60, 70))
    draw.text((60, 105), "Invoice No: INV-2026-088 | Date: 18 Feb 2026", fill=(60, 60, 70))

    # Billed to
    draw.text((60, 170), "Billed To: Orion Interiors", fill=(20, 20, 30))
    draw.text((60, 195), "Reference Agreement: OS-NW-2026-004", fill=(70, 70, 80))

    # Table Header
    draw.rectangle([(60, 240), (940, 280)], fill=(220, 225, 235))
    draw.text((80, 252), "Item Description", fill=(20, 20, 30))
    draw.text((500, 252), "Qty", fill=(20, 20, 30))
    draw.text((650, 252), "Rate (INR)", fill=(20, 20, 30))
    draw.text((800, 252), "Amount (INR)", fill=(20, 20, 30))

    # Items
    draw.text((80, 300), "Commercial Executive Desks & Mesh Chairs", fill=(30, 30, 30))
    draw.text((510, 300), "Lot", fill=(30, 30, 30))
    draw.text((650, 300), "5,00,000", fill=(30, 30, 30))
    draw.text((800, 300), "5,00,000", fill=(30, 30, 30))

    draw.line([(60, 340), (940, 340)], fill=(200, 200, 200))
    draw.text((600, 360), "Subtotal: INR 5,00,000", fill=(30, 30, 30))
    draw.text((600, 390), "CGST (9%): INR 45,000", fill=(30, 30, 30))
    draw.text((600, 420), "SGST (9%): INR 45,000", fill=(30, 30, 30))
    draw.text((600, 460), "Total Invoice Value: INR 5,90,000", fill=(10, 10, 10))

    # Planted conflict!
    draw.rectangle([(580, 510), (940, 600)], fill=(255, 245, 240), outline=(220, 100, 100))
    draw.text((600, 525), "Advance Received: INR 2,00,000", fill=(160, 20, 20))
    draw.text((600, 555), "Balance Due: INR 3,90,000", fill=(160, 20, 20))

    # Notes
    draw.text((60, 650), "Payment Terms: Bank transfer to Northwind Supplies HDFC A/c 50200012345678", fill=(50, 50, 50))
    draw.text((60, 680), "Authorized Signatory: Vikram Sharma", fill=(50, 50, 50))

    # Apply slight rotation/skew to simulate scan
    rotated = img.rotate(0.7, resample=Image.BICUBIC, expand=False, fillcolor=(250, 249, 245))
    rotated.save(img_path, quality=92)
    print("Generated 2_Invoice.jpg")

def generate_vendor_email():
    # 3_Email_Vendor.docx
    doc_path = DEMO_DIR / "3_Email_Vendor.docx"
    doc = Document()
    doc.add_heading("INTERNAL & CLIENT EMAIL ARCHIVE", level=1)
    
    p = doc.add_paragraph()
    p.add_run("From: ").bold = True
    p.add_run("Vikram Sharma <vikram@northwindsupplies.com>\n")
    p.add_run("To: ").bold = True
    p.add_run("Rajesh Mehta <rajesh@orioninteriors.com>\n")
    p.add_run("Date: ").bold = True
    p.add_run("20 Feb 2026, 17:45 IST\n")
    p.add_run("Subject: ").bold = True
    p.add_run("Order Delivery Confirmation - Agreement OS-NW-2026-004\n")

    doc.add_heading("Message Body", level=2)
    doc.add_paragraph(
        "Dear Rajesh,\n\n"
        "We are pleased to inform you that our logistics dispatch team loaded the full consignment of executive desks "
        "and chairs from our warehouse on 12 Feb 2026. The shipment transit was completed and our delivery vehicle "
        "delivered on 20 Feb 2026 to your Mumbai facility loading dock.\n\n"
        "Our site team assisted in unloading all 50 packages. Please release the remaining balance payment of INR 3,90,000 "
        "as per Invoice INV-2026-088.\n\n"
        "Warm regards,\n"
        "Vikram Sharma\n"
        "Director, Northwind Supplies Pvt Ltd"
    )
    doc.save(doc_path)
    print("Generated 3_Email_Vendor.docx")

def generate_client_email():
    # 4_Email_Client.docx
    doc_path = DEMO_DIR / "4_Email_Client.docx"
    doc = Document()
    doc.add_heading("CLIENT COMMUNICATION / DISPUTE NOTICE", level=1)

    p = doc.add_paragraph()
    p.add_run("From: ").bold = True
    p.add_run("Rajesh Mehta <rajesh@orioninteriors.com>\n")
    p.add_run("To: ").bold = True
    p.add_run("Vikram Sharma <vikram@northwindsupplies.com>\n")
    p.add_run("Date: ").bold = True
    p.add_run("02 Mar 2026, 11:20 IST\n")
    p.add_run("Subject: ").bold = True
    p.add_run("DISPUTE: Delayed Delivery & Damaged Desks - OS-NW-2026-004\n")

    doc.add_heading("Formal Notice of Breach", level=2)
    doc.add_paragraph(
        "Dear Vikram,\n\n"
        "We strongly object to your characterization that delivery happened on 20 Feb. In reality, your truck arrived without "
        "unloading personnel and goods were actually received on 28 Feb 2026 at our central facility.\n\n"
        "Furthermore, upon opening the crates on 1 March, our inspection team found severe scratches on 8 executive conference "
        "desks and torn fabric on 4 task chairs. The goods are significantly damaged.\n\n"
        "Under Clause 3.1 of our signed Contract (10 Jan 2026), delivery was due by 15 Feb 2026. Because delivery occurred "
        "on 28 Feb 2026, Orion Interiors is enforcing the late penalty of 1% per week on the order value. Additionally, we require "
        "immediate rectification of the damaged furniture before any further balance is released.\n\n"
        "Sincerely,\n"
        "Rajesh Mehta\n"
        "Managing Director, Orion Interiors"
    )
    doc.save(doc_path)
    print("Generated 4_Email_Client.docx")

def generate_meeting_notes():
    # 5_Meeting_Notes.txt
    txt_path = DEMO_DIR / "5_Meeting_Notes.txt"
    content = """MEETING NOTES - COMMERCIAL REVIEW MEETING
Date: 05 Mar 2026
Attendees: Rajesh Mehta (Orion), Vikram Sharma (Northwind), S. Kothari (Facilitator)
Meeting Location: Orion Interiors HQ, Boardroom B

Discussion Points:
1. Vendor explained transit delays caused by interstate border check delays.
2. Orion pointed out that the contract delivery date was 15 Feb 2026 and buyer actually received goods on 28 Feb 2026.
3. Furniture damage review: Northwind agreed to send carpenters to inspect scratched conference desks on 8 Mar.
4. Financial settlement: Both parties discussed the outstanding invoice and advance dispute.
   Northwind stated: 'payment will be settled soon.'
   No specific settlement amount, deduction calculation, or definitive date was committed.
5. Action Items:
   - Northwind to provide replacement parts.
   - Orion to re-inspect and communicate final penalty deductions.
"""
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(content)
    print("Generated 5_Meeting_Notes.txt")

def generate_payment_receipt():
    # 6_Payment_Receipt.png (low-quality phone photo, blurry)
    png_path = DEMO_DIR / "6_Payment_Receipt.png"
    w, h = 800, 1000
    img = Image.new("RGB", (w, h), color=(240, 240, 235))
    draw = ImageDraw.Draw(img)

    # Phone photo simulation of an e-receipt / bank advice
    draw.rectangle([(50, 60), (750, 920)], fill=(255, 255, 255), outline=(190, 190, 185))
    draw.text((80, 100), "HDFC BANK - NEFT CUSTOMER ACKNOWLEDGEMENT", fill=(30, 40, 80))
    draw.line([(80, 130), (720, 130)], fill=(180, 180, 180))

    draw.text((80, 160), "Transaction Reference: N1102604819230", fill=(50, 50, 50))
    draw.text((80, 200), "Date of Transfer: 11 Jan 2026", fill=(40, 40, 40))
    draw.text((80, 240), "Remitter Name: Orion Interiors", fill=(40, 40, 40))
    draw.text((80, 280), "Beneficiary: Northwind Supplies Pvt Ltd", fill=(40, 40, 40))
    draw.text((80, 320), "Beneficiary Account: 50200012345678", fill=(40, 40, 40))

    draw.rectangle([(80, 370), (720, 470)], fill=(245, 250, 245), outline=(100, 180, 100))
    draw.text((100, 395), "Amount Paid: INR 2,50,000.00", fill=(10, 80, 20))
    draw.text((100, 430), "(Rupees Two Lakh Fifty Thousand Only)", fill=(40, 90, 40))

    draw.text((80, 520), "Purpose: 50% Advance for Office Furniture Supply OS-NW-2026-004", fill=(60, 60, 60))
    draw.text((80, 560), "Status: TRANSACTION SUCCESSFUL", fill=(20, 120, 30))

    draw.text((80, 700), "Generated on: 11 Jan 2026 14:22:10 IST", fill=(120, 120, 120))

    # Apply blur to simulate low quality phone camera
    blurred = img.filter(ImageFilter.GaussianBlur(radius=1.6))
    blurred.save(png_path)
    print("Generated 6_Payment_Receipt.png")

if __name__ == "__main__":
    generate_pdf()
    generate_invoice_image()
    generate_vendor_email()
    generate_client_email()
    generate_meeting_notes()
    generate_payment_receipt()
    print("All demo documents generated successfully!")
